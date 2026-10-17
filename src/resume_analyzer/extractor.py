
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_pdf_text(file_path: str | Path) -> str:
    """Extract text from a text-based PDF."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    reader = PdfReader(str(path))

    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported.")

    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    return "\n".join(pages)


def extract_docx_text(file_path: str | Path) -> str:
    """Extract text from DOCX paragraphs and tables."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    document = Document(str(path))
    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    # Resume content can also appear in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_text(file_path: str | Path) -> str:
    """Dispatch extraction based on the file extension."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Resume not found: {path}")

    extension = path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(path)

    if extension == ".docx":
        return extract_docx_text(path)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Please use PDF or DOCX."
    )
