
from docx import Document

from src.resume_analyzer.cleaner import clean_text
from src.resume_analyzer.extractor import extract_text
from src.resume_analyzer.segmenter import segment_sections


def test_clean_text():
    text = "  Python    Developer  \r\n\r\n\r\nSkills: SQL  "
    cleaned = clean_text(text)

    assert cleaned == "Python Developer\n\nSkills: SQL"


def test_segment_sections():
    text = """
    PROFESSIONAL SUMMARY
    Python developer

    TECHNICAL SKILLS
    Python, SQL

    EDUCATION
    B.Tech
    """

    sections = segment_sections(clean_text(text))

    assert sections["summary"] == "Python developer"
    assert sections["skills"] == "Python, SQL"
    assert sections["education"] == "B.Tech"


def test_extract_docx(tmp_path):
    path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("Alex Morgan")
    document.add_paragraph("TECHNICAL SKILLS")
    document.add_paragraph("Python, SQL")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "University"
    table.cell(0, 1).text = "2026"

    document.save(path)

    text = extract_text(path)

    assert "Alex Morgan" in text
    assert "Python, SQL" in text
    assert "University" in text
    assert "2026" in text
